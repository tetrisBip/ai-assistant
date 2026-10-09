# app/services/file_reader.py
"""
Извлечение текста из файлов. OCR убран — картинки идут в vision-модель.
Синхронные функции — вызывать через asyncio.to_thread.
"""
from pathlib import Path

from app.utils.logger import log


TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst",
    ".log", ".conf", ".cfg", ".ini", ".env",
    ".yaml", ".yml", ".json", ".xml", ".toml",
    ".csv", ".tsv",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".vue", ".svelte",
    ".sh", ".bash", ".zsh", ".ps1", ".bat", ".cmd",
    ".sql", ".go", ".rs", ".c", ".h", ".cpp", ".hpp",
    ".java", ".kt", ".swift", ".rb", ".php", ".pl",
    ".html", ".htm", ".css", ".scss", ".less",
    ".dockerfile", ".gitignore", ".editorconfig",
    ".service", ".socket", ".timer", ".target",
}

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}

MAX_CHARS_DEFAULT = 30000


class FileReadError(Exception):
    pass


def read_text_like(path: Path, max_chars: int = MAX_CHARS_DEFAULT) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[...файл обрезан на {max_chars} символах...]"
    return text


def read_pdf(path: Path, max_chars: int = MAX_CHARS_DEFAULT) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as e:
            log.warning(f"PDF page {i} extract failed: {e}")
            pages.append("")

    text = "\n\n".join(pages).strip()
    if not text:
        raise FileReadError(
            "PDF не содержит извлекаемого текста. Возможно, это скан — "
            "пришли его как картинку (скриншот), обработаю через vision."
        )
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[...обрезано на {max_chars} символах...]"
    return text


def read_docx(path: Path, max_chars: int = MAX_CHARS_DEFAULT) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))

    text = "\n".join(parts).strip()
    if not text:
        raise FileReadError("DOCX пустой или не содержит текста")
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[...обрезано...]"
    return text


def read_xlsx(path: Path, max_chars: int = MAX_CHARS_DEFAULT) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"### Лист: {sheet_name}")
        for row in ws.iter_rows(values_only=True):
            if all(c is None for c in row):
                continue
            parts.append(" | ".join("" if c is None else str(c) for c in row))
            if sum(len(p) for p in parts) > max_chars:
                parts.append("[...обрезано...]")
                break
    text = "\n".join(parts).strip()
    if not text:
        raise FileReadError("XLSX пустой")
    return text


def extract_text(path: Path, max_chars: int = MAX_CHARS_DEFAULT) -> str:
    """
    Диспетчер по расширению. Картинки не читает — для них есть vision.
    """
    if not path.exists():
        raise FileReadError(f"файл не найден: {path}")
    if not path.is_file():
        raise FileReadError(f"это не файл: {path}")

    suffix = path.suffix.lower()

    if suffix in IMAGE_SUFFIXES:
        raise FileReadError(
            "это изображение — пришли его как фото в чат, "
            "я обработаю через vision-модель."
        )

    try:
        if suffix == ".pdf":
            return read_pdf(path, max_chars)
        if suffix == ".docx":
            return read_docx(path, max_chars)
        if suffix in {".xlsx", ".xlsm"}:
            return read_xlsx(path, max_chars)
        if suffix in TEXT_SUFFIXES or not suffix:
            return read_text_like(path, max_chars)
    except FileReadError:
        raise
    except Exception as e:
        log.exception(f"extract_text failed for {path}")
        raise FileReadError(f"ошибка чтения {suffix}: {e}")

    raise FileReadError(f"формат {suffix} пока не поддерживается")


def describe_file(path: Path) -> str:
    try:
        size = path.stat().st_size
    except Exception:
        size = 0
    return f"{path.name} ({path.suffix or 'no-ext'}, {size} байт)"